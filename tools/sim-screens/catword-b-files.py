"""The upload starter files of catword lessons 6-10 and a done-right copy of
each, made with python-docx and openpyxl on the host - a STAND-IN for
catword-b-files.ps1 (real Word and Excel in the CAT VM), because on
8 October 2026 from about 17:30 Word's SaveAs2 hung in the VM for every
script (a plain new document too; Excel saved). Run catword-b-files.ps1 again
once Word saves, and copy its files over these. Same documents, same text.

    C:/Python314/python.exe catword-b-files.py
Writes files/catword-b-files/ here.
"""
import os
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'files', 'catword-b-files')
os.makedirs(OUT, exist_ok=True)

NEWS = [
    ('Title', 'Phumlani News'),
    (None, 'Term 3, 2026 - the newsletter of Phumlani Secondary, Soweto'),
    ('Heading 1', 'Our new library'),
    (None, 'The new library opens on Monday. It has forty computers, a printer and more than two thousand books. Ms Naidoo says the computers are for homework and projects during break and after school, and the librarian will help anyone who gets stuck.'),
    ('Heading 1', 'Market day'),
    (None, "The Grade 10 market day raised R4 350 for the matric farewell. Thank you to every stall holder and every parent who came. The cupcakes sold out by eleven o'clock, and the boerewors rolls soon after."),
    ('Heading 1', 'Sport'),
    (None, 'The netball team beat Orlando High 24 to 19 on Saturday, and the soccer team plays in the district semi-final next week. Come and support them at the Dobsonville stadium.'),
    ('Heading 1', 'Dates to remember'),
    (None, 'Reports go home on Friday 25 September. Schools close on 2 October and open again on 13 October.'),
]
ITEMS = [('White bread', '18.00', '42'), ('Brown bread', '17.00', '35'), ('Cupcakes', '12.00', '60'), ('Koeksisters', '8.00', '48')]


def a4(doc, margin=2.54):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(margin)


def field(paragraph, instr, result):
    """A simple field (PAGE, =SUM(ABOVE)) with its last result, as Word stores it."""
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), instr)
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = result
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def newsletter(done):
    doc = Document()
    a4(doc, 2 if done else 2.54)
    for style, text in NEWS:
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        if done and text == 'Sport':
            p.add_run().add_break(WD_BREAK.PAGE)
        p.add_run(text)
    if done:
        doc.sections[0].header.paragraphs[0].text = 'Phumlani News - Term 3'
        fp = doc.sections[0].footer.paragraphs[0]
        fp.alignment = 1
        field(fp, 'PAGE', '1')
    doc.save(os.path.join(OUT, 'Newsletter-done.docx' if done else 'Newsletter.docx'))


def sales_table(done):
    doc = Document()
    a4(doc)
    doc.add_paragraph('Saturday sales', style='Heading 1')
    doc.add_paragraph("Botha's Bakery, Centurion - what we sold on Saturday.")
    rows = [("Botha's Bakery - Saturday sales", '', ''), ('Item', 'Price (R)', 'Number sold')] + ITEMS
    if done:
        rows.append(('Rusks', '25.00', '30'))
    rows.append(('Total', '', ''))
    t = doc.add_table(rows=len(rows), cols=3)
    t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if v:
                t.cell(i, j).text = v
    if done:
        title = t.cell(0, 0).merge(t.cell(0, 2))
        title.text = "Botha's Bakery - Saturday sales"
        field(t.cell(len(rows) - 1, 2).paragraphs[0], '=SUM(ABOVE)', '215')
    doc.add_paragraph()
    doc.save(os.path.join(OUT, 'SaturdaySales-done.docx' if done else 'SaturdaySales.docx'))


ESSAY = [
    'When I finish school I want to become a software developer. I like solving problems, and I want to recieve a bursary to study at university.',
    'My sister Lerato says the the best developers never stop learning. She is definately right: the tools change every year.',
    'I will practise every day, keep my marks high and finish my projects on time. I beleive nothing will stop me.',
]
FIXES = [('recieve', 'receive'), ('the the', 'the'), ('definately', 'definitely'), ('beleive', 'believe')]


def essay(done):
    doc = Document()
    a4(doc)
    doc.add_paragraph('My dream job', style='Title')
    paras = ESSAY
    if done:
        for a, b in FIXES:
            paras = [p.replace(a, b) for p in paras]
    for p in paras:
        doc.add_paragraph(p)
    count = sum(len(p.split()) for p in paras)
    doc.add_paragraph('Words: %d' % count if done else 'Words:')
    doc.save(os.path.join(OUT, 'ThaboEssay-done.docx' if done else 'ThaboEssay.docx'))
    return count


SHEET = [('Item', 'Price (R)', 'Number sold')] + [(a, float(b), int(c)) for a, b, c in ITEMS] + [('Rusks', 25.0, 30)]


def sales_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Saturday'
    for r, row in enumerate(SHEET, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    ws['A7'] = 'Total'
    ws['C7'] = '=SUM(C2:C6)'
    for c in 'ABC':
        ws[c + '1'].font = Font(bold=True)
        ws[c + '7'].font = Font(bold=True)
    for r in range(2, 7):
        ws['B%d' % r].number_format = '0.00'
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 11
    ws.column_dimensions['C'].width = 13
    wb.save(os.path.join(OUT, 'SaturdaySales.xlsx'))


def report(done):
    doc = Document()
    a4(doc)
    doc.add_paragraph('Saturday report', style='Heading 1')
    doc.add_paragraph("Mr Botha asked for Saturday's sales in his weekly report. Here they are:")
    if done:
        rows = [('Item', 'Price (R)', 'Number sold')] + [(a, '%.2f' % b, str(c)) for a, b, c in SHEET[1:]] + [('Total', '', '215')]
        t = doc.add_table(rows=len(rows), cols=3)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                if v:
                    t.cell(i, j).text = v
        doc.add_paragraph('Source: SaturdaySales.xlsx')
    else:
        doc.add_paragraph()
    doc.add_paragraph('Next week we bake twice as many cupcakes.')
    doc.save(os.path.join(OUT, 'SalesReport-done.docx' if done else 'SalesReport.docx'))


for done in (False, True):
    newsletter(done)
    sales_table(done)
    count = essay(done)
    report(done)
sales_sheet()
open(os.path.join(OUT, 'wordcount.txt'), 'w').write(str(count))
print('word count', count)
for f in sorted(os.listdir(OUT)):
    print(' ', f, os.path.getsize(os.path.join(OUT, f)))
